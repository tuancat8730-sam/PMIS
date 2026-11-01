"""Sinh file .docx từ các section đã điền dữ liệu (không chạm DB).

Hàm chỉ nhận dữ liệu đã chuẩn bị và trả về bytes.
Việc lưu file, đánh số bản và audit nằm ở tầng khác.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO

from docx import Document


@dataclass(frozen=True)
class RenderedSection:
    heading: str
    lines: tuple[str, ...]


def render_report_docx(title: str, sections: Sequence[RenderedSection]) -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    for section in sections:
        document.add_heading(section.heading, level=1)
        for line in section.lines:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
